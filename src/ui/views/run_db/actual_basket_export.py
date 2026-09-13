"""ZIP export for the actual purchasing-result tables."""

from __future__ import annotations

from collections import defaultdict
from collections.abc import Iterable, Mapping
from io import BytesIO
import re
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXCEL_COLUMNS = {
    "item_code": "كود الصنف",
    "item_name": "الصنف المطلوب",
    "matched_name": "الاسم المطابق",
    "requested_qty": "الكمية المطلوبة",
    "excel_public_price": "سعر البيع",
    "excel_purchase_price": "سعر شراء Excel",
    "excel_discount_percent": "نسبة الخصم",
    "status": "الحالة",
}

TAWREED_COLUMNS = {
    "item_code": "كود الصنف",
    "item_name": "الصنف المطلوب",
    "matched_name": "الاسم المطابق",
    "requested_qty": "الكمية المطلوبة",
    "ordered_qty": "الكمية المضافة",
    "public_price": "سعر البيع",
    "purchase_price": "سعر الشراء",
    "discount_percent": "نسبة الخصم",
    "status": "الحالة",
}


def build_actual_basket_zip(
    tawreed_rows: Iterable[Mapping], deferred_rows: Iterable[Mapping]
) -> bytes:
    """Build one independent right-to-left workbook per destination group."""
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        for index, (title, rows, columns) in enumerate(
            _export_groups(tawreed_rows, deferred_rows), start=1
        ):
            archive.writestr(
                f"{index:03d}_{safe_filename(title)}.xlsx",
                _build_destination_workbook(title, rows, columns),
            )
    return output.getvalue()


def _export_groups(
    tawreed_rows: Iterable[Mapping], deferred_rows: Iterable[Mapping]
) -> list[tuple[str, list[Mapping], Mapping[str, str]]]:
    groups: list[tuple[str, list[Mapping], Mapping[str, str]]] = []
    for source, rows in _group_rows_by_key(deferred_rows, "excel_target_source"):
        groups.append((_display_excel_source(source), rows, EXCEL_COLUMNS))
    for store, rows in _store_groups(tawreed_rows):
        groups.append((str(store or "مخزن غير معروف"), rows, TAWREED_COLUMNS))
    return groups


def _build_destination_workbook(
    title: str, rows: list[Mapping], columns: Mapping[str, str]
) -> bytes:
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    _append_sheet(workbook, title, rows, columns, set())

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _group_rows_by_key(
    rows: Iterable[Mapping], key: str
) -> list[tuple[object, list[Mapping]]]:
    grouped: defaultdict[object, list[Mapping]] = defaultdict(list)
    for row in rows:
        grouped[row.get(key)].append(row)
    return sorted(grouped.items(), key=lambda group: str(group[0] or ""))


def _store_groups(rows: Iterable[Mapping]) -> list[tuple[object, list[Mapping]]]:
    """Group Tawreed rows by store identity, keeping same-name branches apart."""
    grouped: defaultdict[object, list[Mapping]] = defaultdict(list)
    for row in rows:
        identity = row.get("store_key") or row.get("store_name") or ""
        grouped[identity].append(row)
    return sorted(
        grouped.items(),
        key=lambda group: (
            str(group[1][0].get("store_name") or ""),
            str(group[0] or ""),
        ),
    )


def _display_excel_source(source: object) -> str:
    return str(source or "Excel Target").split("@", maxsplit=1)[0].strip()


def safe_filename(destination_name: str) -> str:
    name = re.sub(
        r'[<>:"/\\|?*\x00-\x1f]', "_", str(destination_name)
    ).strip(" .")[:110]
    return "actual_basket_" + (name or "unnamed")


def _append_sheet(
    workbook: Workbook,
    title: str,
    rows: list[Mapping],
    columns: Mapping[str, str],
    used_names: set[str],
) -> None:
    sheet = workbook.create_sheet(_sheet_name(title, used_names))
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"
    sheet.append(list(columns.values()))
    for row in rows:
        sheet.append([row.get(key) for key in columns])
    _format_sheet(sheet)


def _sheet_name(sheet_title: str, used_names: set[str]) -> str:
    base = re.sub(r'[:\\/*?\[\]]', "_", sheet_title).strip(" .") or "بدون اسم"
    base = base[:31]
    candidate = base
    suffix = 2
    while candidate.casefold() in used_names:
        suffix_text = f" ({suffix})"
        candidate = base[: 31 - len(suffix_text)] + suffix_text
        suffix += 1
    used_names.add(candidate.casefold())
    return candidate


def _format_sheet(sheet) -> None:
    for cells in sheet:
        for cell in cells:
            cell.font = Font(name="Arial", size=11)
            cell.alignment = Alignment(vertical="center")
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor="24566E")
    for index, width in enumerate((18, 38, 38, 16, 16, 18, 16, 20, 16), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.auto_filter.ref = sheet.dimensions


__all__ = ["build_actual_basket_zip", "safe_filename"]
