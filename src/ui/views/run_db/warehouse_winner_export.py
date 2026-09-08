"""Shared warehouse grouping and Excel columns for display and downloads."""

from __future__ import annotations

from collections import defaultdict
from io import BytesIO
import re
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = {
    "item_code": "كود الصنف", "item_name": "الصنف المطلوب",
    "product_name": "اسم الصنف لدى المورد", "requested_qty": "الكمية المطلوبة",
    "available_qty": "الكمية المتاحة", "public_price": "سعر الجمهور",
    "discount_percent": "الخصم (%)", "purchase_price": "سعر الشراء", "currency": "العملة",
}


def warehouse_groups(rows: list[dict]) -> list[tuple[tuple, list[dict]]]:
    """Group by source and store identity, retaining distinct namesakes."""
    groups = defaultdict(list)
    for row in rows:
        groups[(row["source_kind"], row["store_key"])].append(row)
    return sorted(groups.items(), key=lambda group: (
        group[0][0], group[1][0]["store_name"], group[0][1],
    ))


def warehouse_frame(rows: list[dict]) -> pd.DataFrame:
    """The UI and workbook use precisely the same columns and values."""
    return pd.DataFrame(rows).reindex(columns=list(COLUMNS)).rename(columns=COLUMNS)


def safe_filename(value: str) -> str:
    """A safe component for Windows and ZIP paths, including reserved names."""
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")[:110]
    return "warehouse_" + (name or "unnamed")


def build_warehouse_zip(rows: list[dict]) -> bytes:
    """Export one text-safe Arabic workbook per winning warehouse."""
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        for index, (_, group) in enumerate(warehouse_groups(rows), start=1):
            name = group[0]["store_name"]
            filename = safe_filename(name.removesuffix(".xlsx"))
            archive.writestr(f"{index:03d}_{filename}.xlsx", _workbook(group))
    return output.getvalue()


def _workbook(rows: list[dict]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "الأصناف الفائزة"
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"
    sheet.append(list(COLUMNS.values()))
    for row in rows:
        sheet.append([row.get(key) for key in COLUMNS])
    _format_sheet(sheet)
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def _format_sheet(sheet) -> None:
    for cells in sheet:
        for cell in cells:
            cell.font = Font(name="Arial", size=11)
            cell.alignment = Alignment(vertical="center")
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor="24566E")
    for index, width in enumerate((20, 42, 42, 20, 20, 18, 18, 18, 12), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.auto_filter.ref = sheet.dimensions
