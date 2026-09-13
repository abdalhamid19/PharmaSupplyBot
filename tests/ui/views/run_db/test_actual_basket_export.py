from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook

from src.ui.views.run_db.actual_basket_export import (
    EXCEL_COLUMNS,
    TAWREED_COLUMNS,
    build_actual_basket_zip,
)


def test_export_download_has_independent_workbook_per_branch() -> None:
    tawreed_rows = [
        {
            "store_name": "../Baraka.xlsx",
            "item_code": "=T1",
            "item_name": "Tawreed item",
            "matched_name": "Matched tawreed",
            "requested_qty": 2,
            "ordered_qty": 2,
            "public_price": 15,
            "purchase_price": 10,
            "discount_percent": 33.33,
            "status": "added-to-cart",
        }
    ]
    deferred_rows = [
        {
            "excel_target_source": "Qaisar branch@qaisar.xlsx",
            "item_code": "E1",
            "item_name": "Excel item",
            "matched_name": "Matched excel",
            "requested_qty": 1,
            "excel_public_price": 12,
            "excel_purchase_price": 9,
            "excel_discount_percent": 25,
            "status": "deferred-to-excel-target",
        }
    ]

    with ZipFile(BytesIO(build_actual_basket_zip(tawreed_rows, deferred_rows))) as archive:
        names = archive.namelist()

        assert len(names) == 2
        assert len(set(names)) == 2
        assert all("/" not in name and "\\" not in name for name in names)

        excel_book = load_workbook(BytesIO(archive.read(names[0])), data_only=True)
        tawreed_book = load_workbook(BytesIO(archive.read(names[1])), data_only=True)

        assert excel_book.sheetnames == ["Qaisar branch"]
        assert tawreed_book.sheetnames == ["_Baraka.xlsx"]
        assert [cell.value for cell in excel_book.active[1]] == list(EXCEL_COLUMNS.values())
        assert [cell.value for cell in tawreed_book.active[1]] == list(TAWREED_COLUMNS.values())
        assert excel_book.active["A2"].value == "E1"
        assert tawreed_book.active["A2"].value == "=T1"
        assert tawreed_book.active["A2"].data_type == "s"
        assert excel_book.active.max_row == 2
        assert tawreed_book.active.max_row == 2
        assert excel_book.active.sheet_view.rightToLeft
        assert tawreed_book.active.sheet_view.rightToLeft

        excel_book.close()
        tawreed_book.close()


def test_export_keeps_same_named_branches_in_separate_files() -> None:
    tawreed_rows = [
        {
            "store_name": "فرع متكرر", "store_key": store_key,
            "item_code": store_key, "item_name": "Item",
            "matched_name": "Matched", "requested_qty": 1, "ordered_qty": 1,
            "public_price": 10, "purchase_price": 8,
            "discount_percent": 20, "status": "added-to-cart",
        }
        for store_key in ("branch-a", "branch-b")
    ]

    with ZipFile(BytesIO(build_actual_basket_zip(tawreed_rows, []))) as archive:
        names = archive.namelist()
        assert len(names) == 2
        for name, store_key in zip(names, ("branch-a", "branch-b")):
            workbook = load_workbook(BytesIO(archive.read(name)), data_only=True)
            assert workbook.active[2][0].value == store_key
            workbook.close()
